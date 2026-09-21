"""Builds a deliberately off-brand deck with known, seeded defects for eval 3."""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from PIL import Image, ImageDraw

W, H = Inches(13.333), Inches(7.5)
prs = Presentation(); prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]
WRONG_BLUE = RGBColor(0x00, 0x30, 0x87)     # the "Duke blue" people guess; real navy is 012169
RED, GREEN = RGBColor(0xFF, 0x00, 0x00), RGBColor(0x00, 0xB0, 0x50)

def box(s, x, y, w, h, txt, size, color, font, bold=False, fill=None):
    if fill is not None:
        r = s.shapes.add_shape(1, x, y, w, h); r.fill.solid(); r.fill.fore_color.rgb = fill; r.line.fill.background()
        tf = r.text_frame
    else:
        tf = s.shapes.add_textbox(x, y, w, h).text_frame
    tf.word_wrap = True
    for i, line in enumerate(txt if isinstance(txt, list) else [txt]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run(); r.text = line; r.font.size = Pt(size); r.font.color.rgb = color; r.font.name = font; r.font.bold = bold

img = Image.new("RGB", (900, 500), "white"); d = ImageDraw.Draw(img)
for i, v in enumerate([310, 342, 371, 398]):
    d.rectangle([120 + i * 180, 450 - v, 240 + i * 180, 450], fill=(0, 48, 135))
img.save("premium_chart.png")

s = prs.slides.add_slide(BLANK)                                   # 1 title: wrong blue, fake typed wordmark, Comic Sans
box(s, 0, 0, W, H, "", 10, WRONG_BLUE, "Calibri", fill=WRONG_BLUE)
box(s, Inches(0.6), Inches(0.5), Inches(3), Inches(1), "DUKE", 48, RGBColor(255, 255, 255), "Times New Roman", bold=True)
box(s, Inches(0.8), Inches(2.6), Inches(11), Inches(1.5), "Open Enrollment 2027", 54, RGBColor(255, 255, 255), "Comic Sans MS", bold=True)
box(s, Inches(0.8), Inches(4.2), Inches(11), Inches(1), "Duke Human Resources  |  Benefits", 24, RGBColor(0xB5, 0xB5, 0xB5), "Calibri")

s = prs.slides.add_slide(BLANK)                                   # 2 low contrast body: persimmon + light gray on white
box(s, Inches(0.6), Inches(0.4), Inches(12), Inches(1), "Key dates", 36, WRONG_BLUE, "Comic Sans MS", bold=True)
box(s, Inches(0.8), Inches(1.7), Inches(11.5), Inches(4), ["• Enrollment opens October 19", "• Enrollment closes October 30 at 6 p.m.",
    "• Changes take effect January 1, 2027"], 24, RGBColor(0xE8, 0x99, 0x23), "Calibri")
box(s, Inches(0.8), Inches(6.2), Inches(11.5), Inches(0.6), "If you take no action, your current elections roll over, except reimbursement accounts.", 16, RGBColor(0xB5, 0xB5, 0xB5), "Calibri")

s = prs.slides.add_slide(BLANK)                                   # 3 off-palette red/green boxes
box(s, Inches(0.6), Inches(0.4), Inches(12), Inches(1), "What's changing", 36, WRONG_BLUE, "Comic Sans MS", bold=True)
box(s, Inches(0.8), Inches(1.8), Inches(5.6), Inches(3.5), ["NEW", "Vision plan adds a second provider network"], 22, RGBColor(255, 255, 255), "Calibri", True, GREEN)
box(s, Inches(6.9), Inches(1.8), Inches(5.6), Inches(3.5), ["ENDING", "Legacy dental option closes to new members"], 22, RGBColor(255, 255, 255), "Calibri", True, RED)

s = prs.slides.add_slide(BLANK)                                   # 4 picture with no alt text
box(s, Inches(0.6), Inches(0.4), Inches(12), Inches(1), "Monthly premiums, 2024 to 2027", 36, WRONG_BLUE, "Comic Sans MS", bold=True)
s.shapes.add_picture("premium_chart.png", Inches(2.2), Inches(1.6), height=Inches(5))

s = prs.slides.add_slide(BLANK)                                   # 5 closing: copper text on wrong-blue band (fails contrast)
box(s, 0, 0, W, Inches(2.2), "", 10, WRONG_BLUE, "Calibri", fill=WRONG_BLUE)
box(s, Inches(0.8), Inches(0.6), Inches(11), Inches(1.2), "Questions? hr.duke.edu/enrollment", 36, RGBColor(0xC8, 0x4E, 0x00), "Comic Sans MS", bold=True)
box(s, Inches(0.8), Inches(3), Inches(11), Inches(2), ["Call (919) 684-5600", "Monday to Friday, 8 a.m. to 5 p.m."], 24, RGBColor(0x26, 0x26, 0x26), "Calibri")
prs.save("open-enrollment-2027.pptx"); print("ok")
